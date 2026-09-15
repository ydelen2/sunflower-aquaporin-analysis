# -*- coding: utf-8 -*-
"""
Build Supplementary_Tables.xlsx (Additional file 1, Tables S1 to S11) from
the pipeline outputs under results/.

usage: python build_supplementary_tables.py
"""
import os, re, sys, csv, collections
import openpyxl
from openpyxl.styles import Font
import figlib as L

OUT = os.path.join(L.ROOT, "tables_and_figures", "supplementary_tables", "Supplementary_Tables.xlsx")
P = lambda *p: os.path.join(L.RESULTS, *p)
s1 = L.load_table_s1()
sub = L.load_subfamily_by_protein()
p2g = L.load_protein_gene()
prods = L.load_products()

def tsv(path):
    with open(path, encoding="utf-8") as fh:
        rows = list(csv.reader(fh, delimiter="\t"))
    return rows[0], rows[1:]

def num(x):
    try:
        if x in ("NA", "", None):
            return x
        f = float(x)
        return int(f) if f.is_integer() and "." not in str(x) else f
    except ValueError:
        return x

wb = openpyxl.Workbook(); wb.remove(wb.active)

def sheet(name, title, header, rows, note=None):
    ws = wb.create_sheet(name)
    ws.cell(1, 1, title).font = Font(bold=True)
    if note:
        ws.cell(2, 1, note).font = Font(italic=True)
    for j, h in enumerate(header, 1):
        ws.cell(3, j, h).font = Font(bold=True)
    for i, r in enumerate(rows, 4):
        for j, v in enumerate(r, 1):
            ws.cell(i, j, num(v))
    ws.freeze_panes = "A4"
    return ws

# S1 gene inventory
h, r = tsv(P("09_figures", "supplementary", "Table_S1_aquaporin_genes.tsv"))
sheet("S1_Aquaporin_Genes", "Table S1. The %d sunflower aquaporin genes: subfamily by phylogenetic placement (subfamily) and by RefSeq product name (subfamily_from_name), tree support of the placement, position, gene structure of the representative mRNA and encoded isoforms." % len(r), h, r)

# S2 rejected candidates
h, r = tsv(P("02_gene_family", "03_domain_verification_v2", "failed_verification.tsv"))
rej = [x for x in r if x[1] == "yes"]
h2 = ["protein_id", "product", "length_aa", "tm_helices_DeepTMHMM", "reason_for_exclusion"]
r2 = []
tm = {x[0]: x for x in tsv(P("02_gene_family", "03_domain_verification_v2", "tm_helix_predictions.tsv"))[1]}
for x in rej:
    pid = x[0]
    r2.append([pid, prods.get(pid, "").replace(" [Helianthus annuus]", ""), tm.get(pid, ["", "", "", ""])[3], x[2], "fewer than four transmembrane helices predicted"])
sheet("S2_Excluded_Candidates", "Table S2. Candidates with a PF00230 domain that were excluded because DeepTMHMM predicted fewer than four transmembrane helices.", h2, r2)

# S5 physicochemical and structural features
h, r = tsv(P("02_gene_family", "aquaporin_characterization.tsv"))
gi = h.index("protein_id")
h = h[:1] + ["gene_id", "subfamily_tree"] + h[1:]
r = [[x[0], p2g.get(x[0], {}).get("gene_id", ""), sub.get(x[0], "")] + x[1:] for x in r]
r.sort(key=lambda x: (x[1], x[0]))
sheet("S5_Physicochemical", "Table S5. Physicochemical properties, NPA motifs and ar/R selectivity filter residues of the %d aquaporin protein isoforms." % len(r), h, r)

# S7a duplication, S7c Ka/Ks (one source table)
h, r = tsv(P("02_gene_family", "09_duplication_kaks_v2", "duplication_kaks.tsv"))
sheet("S7a_Duplication", "Table S7a. Paralogous aquaporin gene pairs (protein identity >= 70%%, query coverage >= 70%%; one pair per gene pair) and their duplication class (tandem <= 200 kb, proximal 200 kb to 1 Mb on the same chromosome, segmental otherwise); n = %d." % len(r),
      [c for c in h if c not in ("aligned_codons", "Ka", "Ks", "Ka_Ks", "selection", "note")],
      [[v for c, v in zip(h, x) if c not in ("aligned_codons", "Ka", "Ks", "Ka_Ks", "selection", "note")] for x in r])
sheet("S7c_Ka_Ks", "Table S7c. Ka, Ks and Ka/Ks (Nei-Gojobori with Jukes-Cantor correction) for the same pairs; NA where fewer than 30 codons aligned or synonymous sites were saturated.",
      ["gene1", "gene2", "protein1", "protein2", "type", "aligned_codons", "Ka", "Ks", "Ka_Ks", "selection", "note"],
      [[x[h.index(c)] for c in ["gene1", "gene2", "protein1", "protein2", "type", "aligned_codons", "Ka", "Ks", "Ka_Ks", "selection", "note"]] for x in r])

# S7b synteny
rows = []
for comp, fn, label in (("Ha_self", "Ha_self_intra_aqp_synteny.tsv", "Intra-genomic"), ("Ha_Ls", "Ha_Ls_inter_aqp_synteny.tsv", "Interspecies"), ("Ha_At", "Ha_At_inter_aqp_synteny.tsv", "Interspecies")):
    h, r = tsv(P("07_synteny", "results_v2", fn))
    for x in r:
        d = dict(zip(h, x))
        if label == "Interspecies" and d["species1"] == d["species2"]:
            continue
        rows.append([comp, label, L.chr_label(d["chr1"]) or d["chr1"], L.chr_label(d["chr2"]) or d["chr2"], d["species1"], d["gene1"], d["prot1"], d["species2"], d["gene2"], d["prot2"]])
sheet("S7b_Synteny", "Table S7b. Collinear aquaporin gene pairs from MCScanX: intra-genomic pairs (both genes in the inventory) and interspecies pairs with lettuce and Arabidopsis (sunflower member in the inventory); n = %d." % len(rows),
      ["Comparison", "Type", "Chr1", "Chr2", "Species1", "Gene1", "Protein1", "Species2", "Gene2", "Protein2"], rows)

# S8 DEG matrix
h, r = tsv(P("09_figures", "supplementary", "Table_S4_aqp_deg_all_contrasts.tsv"))
sheet("S8_AQP_DEGs", "Table S8. DESeq2 log2 fold change and adjusted P value of every aquaporin gene in the 16 stress contrasts (NA: no estimate).", h, r)

# S9, S10 WGCNA
h, r = tsv(P("04_expression", "results", "wgcna", "aquaporin_module_membership.tsv"))
sheet("S9_WGCNA_Membership", "Table S9. WGCNA module assignment of the aquaporin genes present in the PRJNA869183 expression matrix (module 0 = unassigned) and module membership (kME) values.", h, r)
h, r = tsv(P("04_expression", "results", "wgcna", "module_trait_table.tsv"))
traits = sorted({x[1] for x in r}); mods = sorted({x[0] for x in r}, key=lambda m: int(m[2:]))
table = {(x[0], x[1]): (x[2], x[3]) for x in r}
sheet("S10_Module_Trait", "Table S10. Module-trait correlations (Pearson r and P value) for all WGCNA modules of PRJNA869183.",
      ["module"] + ["r_%s" % t for t in traits] + ["p_%s" % t for t in traits],
      [[m] + [table.get((m, t), ("", ""))[0] for t in traits] + [table.get((m, t), ("", ""))[1] for t in traits] for m in mods])

# S11 GO enrichment
h, r = tsv(P("04_expression", "results", "enrichment", "GO_all_contrasts_v2.tsv"))
n_uni = sum(1 for line in open(P("04_expression", "results", "enrichment", "GO_all_contrasts_v2_universe.txt")) if line.strip())
sheet("S11_GO_Enrichment", "Table S11. GO terms enriched in the DEG set of each of the 16 stress contrasts (over-representation analysis with clusterProfiler against one common universe of {:,} genes; {:,} terms).".format(n_uni, len(r)), h, r)

# S6 MEME motifs (protein)
h, r = tsv(P("05_phylogenetics", "results", "meme_motifs", "motif_summary.tsv"))
sheet("S6_MEME_Motifs", "Table S6. Conserved protein motifs identified by MEME (zoops model, 20 motifs searched, ten retained) in the %d aquaporin isoforms." % len(sub), h, r[:10])

# S4 LOC-TAIR mapping
h, r = tsv(P("09_figures", "supplementary", "Table_S9_loc_tair_mapping.tsv"))
sheet("S4_LOC_TAIR_Mapping", "Table S4. LOC-to-TAIR ortholog mapping used for GO enrichment (BLASTp best hit, E-value <= 1e-5; {:,} pairs).".format(len(r)), h, r)

# S3 DESeq2 design rationale; the PRJNA492303 row reflects the 96-run analysis
h, r = tsv(P("09_figures", "supplementary", "Table_S10_design_rationale.tsv"))
s = tsv(P("04_expression", "results", "deseq2", "PRJNA492303_fixed", "DEG_summary.tsv"))[1][0]
for x in r:
    if x[0] == "PRJNA492303":
        x[1] = "~tissue + condition (46 of the 96 runs)"
        x[2] = "~genotype + tissue + age + condition (all 96 runs)"
        x[3] = "The submitted analysis used an arbitrary block of 46 consecutive runs in which HA351 occurred only among controls; the experiment is balanced (2 genotypes x 2 treatments x 2 tissues x 3 ages x 4 replicates)"
        x[4] = "Flooding 1,671 -> %s genome-wide DEGs (%s up, %s down); aquaporin DEGs 5 -> 3" % (s[3], s[4], s[5])
sheet("S3_Design_Rationale", "Table S3. DESeq2 design rationale and correction summary for all six BioProjects.", h, r)

wb._sheets.sort(key=lambda ws: (int(re.match(r"S(\d+)", ws.title).group(1)), ws.title))
wb.properties.creator = "Yavuz Delen"; wb.properties.lastModifiedBy = "Yavuz Delen"
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT)
for ws in wb.worksheets:
    print("  %-24s rows %d" % (ws.title, ws.max_row))
